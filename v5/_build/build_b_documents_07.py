"""Reuse the verified 06 typography for synchronized revision 07 documents."""
from pathlib import Path
import json
import hashlib
import argparse
import build_d_documents_06 as common
from docx.shared import Pt
ROOT=Path(__file__).resolve().parents[1]
JOBS=[('14_B部分口播与练习.md','B部分口播与练习_V5_07'),
      ('15_B部分答辩备课.md','组长B部分答辩备课_V5_07'),
      ('01_自然口播稿.md','全场口播_V5_07'),
      ('04_来源与答辩.md','来源与答辩_V5_07'),
      ('05_课程报告.md','课程报告_V5_07')]

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--replace-draft',action='store_true',help='Rebuild only this task-owned 07 draft')
    parser.add_argument('--b-only',action='store_true')
    args=parser.parse_args()
    receipt=ROOT/'qa/b_integration_07/documents_build.json'
    records=json.loads(receipt.read_text(encoding='utf-8')) if args.b_only and receipt.exists() else []
    for source,stem in JOBS:
        if args.b_only and not source.startswith(('14_','15_')):continue
        target=ROOT/'outputs'/f'{stem}.docx'
        if target.exists() and not args.replace_draft:raise FileExistsError(target)
        common.base.convert(ROOT/source,target)
        if source=='01_自然口播稿.md':common.polish_document(target,source)
        if source.startswith(('14_','15_')):
            doc=common.base.Document(target)
            normal=doc.styles['Normal'].paragraph_format
            normal.line_spacing=1.10 if source.startswith('14_') else 1.13
            normal.space_after=Pt(6)
            for style in ['Heading 1','Heading 2']:
                doc.styles[style].paragraph_format.space_before=Pt(12)
                doc.styles[style].paragraph_format.space_after=Pt(5)
            for index,paragraph in enumerate(doc.paragraphs):
                paragraph.paragraph_format.keep_together=True
                if paragraph.text.startswith(('短答','回答：')):paragraph.paragraph_format.keep_with_next=True
                if paragraph.text.startswith(('指图','动作')) and index:
                    doc.paragraphs[index-1].paragraph_format.keep_with_next=True
            doc.save(target)
        records=[r for r in records if r['source']!=source]
        records.append({'source':source,'sourceSha256':hashlib.sha256((ROOT/source).read_bytes()).hexdigest(),
                        'docx':target.name,'docxSha256':hashlib.sha256(target.read_bytes()).hexdigest()})
    (ROOT/'qa/b_integration_07/documents_build.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
    print('DOCX_07_CREATED',len(records))
if __name__=='__main__':main()
