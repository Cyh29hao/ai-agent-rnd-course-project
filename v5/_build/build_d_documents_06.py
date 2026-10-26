"""Create revision 06 documents without touching previous delivered files."""
from pathlib import Path
import hashlib
import json
import re
import shutil
import sys
import build_documents as base
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[1]
JOBS=[
    ('11_D部分口播与练习.md','D部分口播与练习_V5_06'),
    ('10_D部分答辩备课.md','组长D部分答辩备课_V5_06'),
    ('01_自然口播稿.md','全场口播_V5_06'),
    ('04_来源与答辩.md','来源与答辩_V5_06'),
    ('05_课程报告.md','课程报告_V5_06'),
]

# Internal links remain readable labels in print. Preserve actual source links.
def printable_inline(paragraph,text,bold=False):
    text=re.sub(r'\[([^]]+)\]\((?!https?://)[^)]+\)',r'\1',text)
    def plain(value):
        if not value:return
        run=paragraph.add_run(value.replace('`','').replace('*',''))
        if bold:run.bold=True
    position=0
    for match in re.finditer(r'\*\*(.+?)\*\*|\[([^]]+)\]\((https?://[^)]+)\)',text):
        plain(text[position:match.start()])
        if match.group(1) is not None:
            printable_inline(paragraph,match.group(1),True)
        else:
            relation=paragraph.part.relate_to(match.group(3),'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),relation)
            run=OxmlElement('w:r');properties=OxmlElement('w:rPr')
            fonts=OxmlElement('w:rFonts');fonts.set(qn('w:eastAsia'),'Microsoft YaHei');properties.append(fonts)
            if bold:properties.append(OxmlElement('w:b'))
            run.append(properties);label=OxmlElement('w:t');label.text=match.group(2);run.append(label)
            link.append(run);paragraph._p.append(link)
        position=match.end()
    plain(text[position:])
base.inline=printable_inline

def polish_document(target,source):
    document=base.Document(target)
    if source=='11_D部分口播与练习.md':
        normal=document.styles['Normal'].paragraph_format
        normal.line_spacing=1.10
        normal.space_after=Pt(6)
        for style in ['Heading 1','Heading 2']:
            document.styles[style].paragraph_format.space_before=Pt(12)
            document.styles[style].paragraph_format.space_after=Pt(5)
        for index,paragraph in enumerate(document.paragraphs):
            if paragraph.text.startswith('指图与动作：') and index:
                document.paragraphs[index-1].paragraph_format.keep_with_next=True
                paragraph.paragraph_format.keep_together=True
    if source=='10_D部分答辩备课.md':
        document.styles['Normal'].paragraph_format.line_spacing=1.13
        document.styles['Normal'].paragraph_format.space_after=Pt(6)
        document.paragraphs[0].text='组长D部分答辩备课'
        for paragraph in document.paragraphs:
            paragraph.paragraph_format.keep_together=True
            if paragraph.text.startswith('短答：'):
                paragraph.paragraph_format.keep_with_next=True
    if source=='01_自然口播稿.md':
        document.styles['Normal'].paragraph_format.line_spacing=1.13
        document.styles['Normal'].paragraph_format.space_after=Pt(6)
    document.save(target)

def main():
    records=[]
    only=next((arg.split('=',1)[1].split(',') for arg in sys.argv if arg.startswith('--only=')),None)
    previous={r['source']:r for r in json.loads((ROOT/'qa/d_refinement_06/documents_build.json').read_text(encoding='utf-8'))} if only else {}
    for source,stem in JOBS:
        if only and source[:2] not in only:
            records.append(previous[source])
            continue
        target=ROOT/'outputs'/f'{stem}.docx'
        if target.exists():
            if '--replace-draft' not in sys.argv:
                raise FileExistsError(f'Preserve previous output: {target}')
            backup=ROOT/'_build/d_refinement_06/documents_attempt1'
            backup.mkdir(exist_ok=True)
            if not (backup/target.name).exists():shutil.copy2(target,backup/target.name)
        base.convert(ROOT/source,target)
        polish_document(target,source)
        records.append({'source':source,'sourceSha256':hashlib.sha256((ROOT/source).read_bytes()).hexdigest(),
                        'docx':target.name,'docxSha256':hashlib.sha256(target.read_bytes()).hexdigest()})
    (ROOT/'qa/d_refinement_06/documents_build.json').write_text(
        json.dumps(records,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print('REVISION_06_DOCUMENTS_CREATED',len(records))

if __name__=='__main__':
    main()
