"""Reuse the verified V5 document typography for the synchronized 08 set."""
from pathlib import Path
import argparse,json,hashlib,re
import build_d_documents_06 as common
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
JOBS=[('17_A部分口播与练习.md','A部分口播与练习_V5_10'),('18_A部分答辩备课.md','组长A部分答辩备课_V5_10'),
      ('01_自然口播稿.md','全场口播_V5_10'),]
def main():
    parser=argparse.ArgumentParser();parser.add_argument('--replace-draft',action='store_true');parser.add_argument('--a-only',action='store_true');parser.add_argument('--sources',nargs='+');args=parser.parse_args()
    receipt=ROOT/'qa/inspiration_10/documents_build.json'
    records=json.loads(receipt.read_text(encoding='utf-8')) if (args.a_only or args.sources) and receipt.exists() else []
    for source,stem in JOBS:
        if args.a_only and not source.startswith(('17_','18_')):continue
        if args.sources and not source.startswith(tuple(args.sources)):continue
        target=ROOT/'outputs'/f'{stem}.docx'
        if target.exists() and not args.replace_draft:raise FileExistsError(target)
        common.base.convert(ROOT/source,target)
        if source=='01_自然口播稿.md':common.polish_document(target,source)
        doc=common.base.Document(target)
        if source.startswith(('17_','18_')):
            doc.styles['Normal'].paragraph_format.line_spacing=1.10
            doc.styles['Normal'].paragraph_format.space_after=Pt(6)
            for name in ['Heading 1','Heading 2']:
                doc.styles[name].paragraph_format.space_before=Pt(11)
                doc.styles[name].paragraph_format.space_after=Pt(5)
            for index,p in enumerate(doc.paragraphs):
                p.paragraph_format.keep_together=True
                if p.text.startswith(('短答','回答：')):p.paragraph_format.keep_with_next=True
                if p.text.startswith('指图') and index:doc.paragraphs[index-1].paragraph_format.keep_with_next=True
        # Repeat table headers across every printed continuation page.
        for table in doc.tables:
            props=table.rows[0]._tr.get_or_add_trPr()
            if props.find(qn('w:tblHeader')) is None:props.append(OxmlElement('w:tblHeader'))
        if source=='05_课程报告.md':
            title=doc.paragraphs[0]
            if '：' in title.text:
                left,right=title.text.split('：',1)
                title.text=left+'\n'+right
                title.paragraph_format.line_spacing=1.05
                for run in title.runs:run.font.size=Pt(17)
            in_references=False
            for paragraph in doc.paragraphs:
                if paragraph.text=='参考文献':in_references=True
                elif in_references and re.match(r'^\[\d+\]',paragraph.text):
                    paragraph.paragraph_format.space_before=Pt(0)
                    paragraph.paragraph_format.space_after=Pt(1)
                    paragraph.paragraph_format.line_spacing=1.05
        doc.save(target)
        records=[r for r in records if r['source']!=source]
        records.append({'source':source,'sourceSha256':hashlib.sha256((ROOT/source).read_bytes()).hexdigest(),
                        'docx':target.name,'docxSha256':hashlib.sha256(target.read_bytes()).hexdigest()})
    receipt.write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
    print('DOCX08',len(records))
if __name__=='__main__':main()
