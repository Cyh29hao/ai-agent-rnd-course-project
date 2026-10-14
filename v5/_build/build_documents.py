"""Build printable documents with the bundled python-docx runtime."""
from pathlib import Path
import re,json
from docx import Document
from docx.shared import Cm,Pt,RGBColor,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1]

def setup():
    d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        st=d.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        pp=st.element.get_or_add_pPr()
        for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
        border=OxmlElement('w:pBdr')
        for side in ['top','left','bottom','right','between','bar']:
            el=OxmlElement('w:'+side);el.set(qn('w:val'),'nil');border.append(el)
        pp.append(border)
    normal=d.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.18;normal.paragraph_format.space_after=Pt(7)
    for name,size in [('Title',20),('Heading 1',15),('Heading 2',12),('Heading 3',11.5)]:
        d.styles[name].font.size=Pt(size);d.styles[name].paragraph_format.keep_with_next=True
    f=sec.footer.paragraphs[0];f.alignment=2
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    return d

def inline(para,text):
    # Preserve real source URLs as clickable hyperlinks in Word.
    pat=r'\[([^]]+)\]\((https?://[^)]+)\)'
    pos=0
    for m in re.finditer(pat,text):
        para.add_run(re.sub(r'[*`]','',text[pos:m.start()]))
        rel=para.part.relate_to(m.group(2),'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
        h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
        r=OxmlElement('w:r');p=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:eastAsia'),'Microsoft YaHei');p.append(fonts);r.append(p)
        t=OxmlElement('w:t');t.text=m.group(1);r.append(t);h.append(r);para._p.append(h);pos=m.end()
    para.add_run(re.sub(r'[*`]','',text[pos:]))

def convert(src,out):
    d=setup();rows=src.read_text(encoding='utf-8').splitlines();i=0
    while i<len(rows):
        t=rows[i].strip();i+=1
        if not t:continue
        if t.startswith('|'):
            raw=[t]
            while i<len(rows) and rows[i].lstrip().startswith('|'):raw.append(rows[i].strip());i+=1
            vals=[[c.strip() for c in row.strip('|').split('|')] for row in raw if not re.match(r'^\|[\s:|\-]+\|$',row)]
            tab=d.add_table(rows=len(vals),cols=len(vals[0]));tab.autofit=False
            n=len(vals[0]);widths=([2.8,4.8,5.1,4.2] if n==4 else [3.7,6.6,6.6] if n==3 else [16.9/n]*n)
            for ri,v in enumerate(vals):
                rp=tab.rows[ri]._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
                if ri==0:rp.append(OxmlElement('w:tblHeader'))
                for ci,text in enumerate(v):
                    cell=tab.cell(ri,ci);cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pp=cell.paragraphs[0];inline(pp,text);pp.paragraph_format.space_after=Pt(5);pp.paragraph_format.space_before=Pt(5)
                    for run in pp.runs:run.font.size=Pt(10);run.bold=ri==0
                    cp=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                    for side in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    cp.append(borders)
                    if ri==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EBEE');cp.append(shade)
            d.add_paragraph();continue
        h=re.match(r'^(#{1,3})\s+(.*)',t)
        if h:
            depth=len(h.group(1));p=d.add_paragraph(style='Title' if depth==1 else 'Heading '+str(depth-1));inline(p,h.group(2));continue
        p=d.add_paragraph();inline(p,t)
        if src.name=='01_自然口播稿.md' or t.startswith('**['):
            p.paragraph_format.keep_together=True
    d.save(out)

jobs=[('01_自然口播稿.md','四人口播_V5.docx'),('04_来源与答辩.md','来源与答辩_V5.docx'),('05_课程报告.md','课程报告_V5.docx')]
for source,target in jobs:convert(ROOT/source,ROOT/'outputs'/target)
print('DOCX_CREATED',len(jobs))
