"""Verify current exported deliverables; human acceptance stays explicitly pending."""
from pathlib import Path
import json,re,hashlib,zipfile
from datetime import datetime,timezone
import fitz
from pptx import Presentation
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
data=json.loads((ROOT/'_build'/'slide_content.json').read_text(encoding='utf-8'))
deck=Presentation(ROOT/'01_正式汇报.pptx')
checks=[]
def record(name,ok,detail): checks.append({'check':name,'passed':bool(ok),'detail':detail})
record('main_count_and_time',len(data)==14 and sum(x['seconds'] for x in data)==550,'14 main pages, 550 allocated seconds; this is allocation, not human rehearsal')
record('speaker_allocations', {s:sum(x['seconds'] for x in data if x['speaker']==s) for s in 'ABCD'}==dict(A=150,B=150,C=130,D=120),'A150 B150 C130 D120')
record('backups_exist',len(deck.slides)>=24,f'{len(deck.slides)} slides, {len(deck.slides)-14} backup pages')
sha=lambda path:hashlib.sha256(path.read_bytes()).hexdigest()
export=json.loads((ROOT/'qa/export_manifest.json').read_text(encoding='utf-8'))
render=json.loads((ROOT/'qa/render_manifest.json').read_text(encoding='utf-8'))
record('current_ppt_is_exported_source',sha(ROOT/'01_正式汇报.pptx')==export['source_sha256'],'Source hash captured before PowerPoint export; not inferred from page count')
record('current_full_pdf_is_exported_pdf',sha(ROOT/'03_完整备份版.pdf')==export['pdf_sha256'],'Exported PDF hash agrees with current file')
record('current_main_pdf_and_png_source',sha(ROOT/'03_完整备份版.pdf')==render['source_pdf_sha256'] and sha(ROOT/'02_主讲版.pdf')==render['main_pdf_sha256'],'Rendering and first 14-page PDF derived from this full PDF')
drift_png=[name for name,digest in render['png_sha256'].items() if sha(ROOT/'qa/slides'/name)!=digest]
record('rendered_png_hashes_current',not drift_png,{'drift':drift_png,'pages':len(render['png_sha256'])})
with fitz.open(ROOT/'02_主讲版.pdf') as main,fitz.open(ROOT/'03_完整备份版.pdf') as full:
    differences=[]
    for i in range(14):
        a=main[i].get_pixmap(matrix=fitz.Matrix(.4,.4),alpha=False).samples
        b=full[i].get_pixmap(matrix=fitz.Matrix(.4,.4),alpha=False).samples
        if hashlib.sha256(a).digest()!=hashlib.sha256(b).digest():differences.append(i+1)
record('main_pages_render_identically_to_full_prefix',not differences,{'mismatch_pages':differences})
missing=[]
for i,row in enumerate(data):
 note=deck.slides[i].notes_slide.notes_text_frame.text
 if row['script'] not in note: missing.append(i+1)
record('notes_match_current_script',not missing,{'mismatching_pages':missing})
script_md=(ROOT/'01_四人逐字稿.md').read_text(encoding='utf-8')
record('script_markdown_matches_current_content',all(row['script'] in script_md for row in data),'All 14 spoken passages match the content source and PPT notes')
for name,n in [('02_主讲版.pdf',14),('03_完整备份版.pdf',len(deck.slides))]:
 with fitz.open(ROOT/name) as pdf:
  record(name,len(pdf)==n,{'pages':len(pdf),'expected':n})
  bad=[i+1 for i,p in enumerate(pdf) if not p.get_text().strip() or '\ufffd' in p.get_text()]
  record(name+'_text',not bad,{'empty_or_replacement_glyph_pages':bad})
with fitz.open(ROOT/'04_课程报告.pdf') as pdf:
 record('report_pages',len(pdf)>=6,{'pages':len(pdf)})
 report_text='\n'.join(p.get_text() for p in pdf)
 required=['AlphaEvolve','AgentEvolver','Runme','13/16','12/16','个人分工','23.3%','65%']
 record('report_key_content',all(x in report_text for x in required),{'missing':[x for x in required if x not in report_text]})
doc=Document(ROOT/'04_课程报告.docx')
record('docx_chinese',any('智能体在大模型产业研发中的应用' in p.text for p in doc.paragraphs),'Read DOCX stored text with python-docx')
handbook=json.loads((ROOT/'qa/speaker_handbook/build_manifest.json').read_text(encoding='utf-8'))
handbook_ok=handbook['visual_review']=='pass' and all(sha(ROOT/name)==digest for name,digest in handbook['sources'].items()) and sha(ROOT/'12_四人讲稿与上场卡.docx')==handbook['docx_sha256'] and sha(ROOT/'12_四人讲稿与上场卡.pdf')==handbook['pdf_sha256']
with fitz.open(ROOT/'12_四人讲稿与上场卡.pdf') as handbook_pdf:
    normalized=re.sub(r'\s+','','\n'.join(page.get_text() for page in handbook_pdf))
    handbook_ok=handbook_ok and len(handbook_pdf)==12 and all(re.sub(r'\s+','',row['script']) in normalized for row in data)
record('printable_handbook_current_and_reviewed',handbook_ok,'12 pages reviewed; sources and exported bytes agree; all 14 spoken passages preserved. Human rehearsal pending')
review=(ROOT/'qa/报告最终复审.md').read_text(encoding='utf-8').lower()
incremental=ROOT/'qa/最终边界与版式增量复审.md'
if incremental.exists(): review+='\n'+incremental.read_text(encoding='utf-8').lower()
report_drift=[name for name in ['04_课程报告.md','04_课程报告.docx','04_课程报告.pdf'] if sha(ROOT/name) not in review]
record('report_matches_visually_reviewed_hashes',not report_drift,{'changed_since_latest_report_review':report_drift})
group=json.loads((ROOT/'_build/group_config.json').read_text(encoding='utf-8'))
cover='\n'.join(q.text for q in deck.slides[0].shapes if q.has_text_frame)
record('member_names_synced',all(row['name'] in cover and row['name'] in report_text for row in group['members']),{'members':[row['name'] for row in group['members']],'roles_not_assigned':all(row['speaker_role'] is None for row in group['members'])})
record('no_group_number_field_on_cover','小组号' not in cover,'User confirmed there is no group number')
with zipfile.ZipFile(ROOT/'01_正式汇报.pptx') as z:
 movies=[x for x in z.namelist() if x.startswith('ppt/media/') and x.endswith('.mp4')]
 record('embedded_mp4',len(movies)==2,{'media_parts':movies})
 embedded={hashlib.sha256(z.read(name)).hexdigest() for name in movies}
 original={sha(ROOT/'media'/name) for name in ['anthropic_cycle.mp4','smoke_log_replay.mp4']}
 record('embedded_movies_match_decoded_originals',embedded==original,'Embedded byte hashes match the two original decoded videos')
scores=json.loads((ROOT/'practice/runs/demo/score.json').read_text(encoding='utf-8'))
record('demo_actual_score',all(v['total']==16 and v['valid_outputs']==16 for v in scores['variants'].values()),{n:{k:v[k] for k in ['total','valid_outputs','label_correct','false_release','location_correct']} for n,v in scores['variants'].items()})
freeze=json.loads((ROOT/'practice/runs/demo/freeze.json').read_text(encoding='utf-8'))
drift=[name for name,digest in freeze['sha256'].items() if hashlib.sha256((ROOT/'practice'/name).read_bytes()).hexdigest()!=digest]
record('frozen_experiment_files_unchanged',not drift,{'drift':drift})
trace=json.loads((ROOT/'practice/runs/demo/public_trace_manifest.json').read_text(encoding='utf-8'))
call_keys={'event_timestamp','call_id','tool','public_input','nested_tool_references','original_event_sha256'}
receipt_keys={'event_timestamp','call_id','original_output_sha256','output_characters','original_event_sha256','content_policy'}
for variant,meta in trace['variants'].items():
    base=ROOT/'practice/runs/demo'/variant
    calls=[json.loads(line) for line in (base/'public_tool_calls.original.jsonl').read_text(encoding='utf-8').splitlines() if line]
    receipts=[json.loads(line) for line in (base/'public_tool_output_receipts.jsonl').read_text(encoding='utf-8').splitlines() if line]
    safe=all(set(row)<=call_keys for row in calls) and all(set(row)<=receipt_keys for row in receipts)
    matched={row['call_id'] for row in calls}=={row['call_id'] for row in receipts}
    counts=len(calls)==meta['public_call_records'] and len(receipts)==meta['tool_receipts'] and sum(row['tool'] in ['exec','functions.exec'] for row in calls)==meta['exec_wrapper_calls'] and sum(len(row['nested_tool_references']) for row in calls)==meta['nested_reference_count']
    record('original_public_trace_'+variant,safe and matched and counts and meta['model_identifier_from_session_metadata']=='gpt-6-sol',{'calls':len(calls),'receipts':len(receipts),'whitelist_and_call_ids_verified':safe and matched,'count_boundary':'Not billing or efficiency measurements'})
for name in ['guide_qa_only','guide_agent_qa']:
    path=ROOT/f'practice/runs/students/{name}/run_manifest.json'
    meta=json.loads(path.read_text(encoding='utf-8'))
    record('guide_record_not_student_identity_'+name,meta.get('executor_name') not in [x['name'] for x in group['members']],{'record':name,'executor':meta.get('executor_name'),'status':meta.get('status')})
for name in ['smoke','anthropic']:
 decode=json.loads((ROOT/f'qa/media_decode/{name}/decode.json').read_text(encoding='utf-8-sig'))
 record('actual_video_decode_'+name,decode['decoded'] and decode['width']==1920 and decode['duration']==(18 if name=='smoke' else 10),decode)
pending=['Four students personally operate and confirm contribution','Human source-location audit and gold-label arbitration','Four-person continuous timed rehearsal including media and handoffs','Offline playback on classroom presentation computer','Latest course deadline and submission format']
output={'verified_at':datetime.now(timezone.utc).isoformat(),'machine_checks':checks,'machine_pass':all(x['passed'] for x in checks),'human_pending':pending,'overall_acceptance':'NOT COMPLETE: machine verification does not replace visual, factual or human acceptance'}
(ROOT/'qa/package_verification.json').write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'machine_pass':output['machine_pass'],'checks':len(checks),'failures':[item for item in checks if not item['passed']],'report':'qa/package_verification.json','overall_acceptance':output['overall_acceptance']},ensure_ascii=True))
raise SystemExit(0 if output['machine_pass'] else 1)
