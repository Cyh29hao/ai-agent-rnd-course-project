"""Build and visually inspect the Chinese course report from UTF-8 Markdown."""
from __future__ import annotations

import re
from pathlib import Path

import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "04_课程报告.md"
DOCX = ROOT / "04_课程报告.docx"
PDF = ROOT / "04_课程报告.pdf"
QA = ROOT / "qa" / "report"
FONT = "等线"
FONT_FILE = Path("C:/Windows/Fonts/msyh.ttc")


def make_figures():
    QA.mkdir(parents=True, exist_ok=True)
    face = lambda n: ImageFont.truetype(str(FONT_FILE), n)
    navy, blue, pale, gray = "#16364c", "#2e7096", "#eaf2f6", "#5d6971"

    im = Image.new("RGB", (1600, 340), "white")
    d = ImageDraw.Draw(im)
    stages = [
        ("共享文献综述", "四个文献 Agent"),
        ("五个并行研究 Agent", "读综述与排行榜"),
        ("候选方法与代码", "结果前冻结 mini-paper"),
        ("后训练目标模型", "提交训练后权重"),
        ("分离评测器", "反馈到论坛与记忆"),
    ]
    for i, (top, bottom) in enumerate(stages):
        x = 22 + i * 318
        d.rounded_rectangle((x, 67, x + 260, 230), 18, fill=pale, outline=blue, width=3)
        d.text((x + 18, 94), top, font=face(25), fill=navy)
        d.text((x + 18, 154), bottom, font=face(21), fill=gray)
        if i < 4:
            d.line((x + 266, 148, x + 305, 148), fill=blue, width=5)
            d.polygon([(x + 305, 148), (x + 288, 137), (x + 288, 159)], fill=blue)
    d.text((25, 267), "权限与评测边界：代码审核、能力退化闸门、留出数据隔离；人设任务和资源上限", font=face(22), fill=gray)
    im.save(QA / "figure_aar_cycle.png")

    im = Image.new("RGB", (1600, 450), "white")
    d = ImageDraw.Draw(im)
    d.text((35, 18), "谄媚任务：同一前150个已评分方法", font=face(30), fill=navy)
    d.text((35, 65), "纵向为不同方法限制；横向为爬坡基准的安全改善空间关闭比例", font=face(21), fill=gray)
    bars = [("不受限", 23.3, "#2e7096"), ("公开数据＋纯监督", 6.0, "#6f9cb6"), ("自由数据＋纯监督", 2.9, "#a8bdc9")]
    x0, scale = 410, 40
    for i, (name, val, color) in enumerate(bars):
        y = 123 + i * 92
        d.text((35, y + 13), name, font=face(23), fill=navy)
        d.rounded_rectangle((x0, y, x0 + int(val * scale), y + 51), 8, fill=color)
        d.text((x0 + int(val * scale) + 18, y + 7), f"{val:.1f}%", font=face(28), fill=navy)
    d.text((35, 410), "注：KL分支在此预算下的端点未从文字核实；完整运行峰值属于另一比较。来源：S2 图20。", font=face(18), fill=gray)
    im.save(QA / "figure_fixed_budget.png")


def font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn("w:" + attr), FONT)


def add_inline(paragraph, line, size=10):
    # Simple inline bold; leave citation identifiers as readable report text.
    for i, part in enumerate(re.split(r"(\*\*.*?\*\*)", line)):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        text = part[2:-2] if bold else part
        r = paragraph.add_run(text)
        font(r, size, bold)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_cell_text(cell, value, header=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(value.strip())
    font(r, 8.2, header)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        shade(cell, "E6EDF2")


def parse_rows(lines):
    rows = []
    for line in lines:
        values = [x.strip() for x in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", v) for v in values):
            continue
        rows.append(values)
    return rows


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            set_cell_text(table.cell(i, j), value, i == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_docx():
    make_figures()
    raw = SOURCE.read_text(encoding="utf-8")
    lines = raw.splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.7), Cm(2.5)
    sec.header_distance, sec.footer_distance = Cm(1.3), Cm(1.3)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    for name, size, before, after in (("Title", 19, 0, 15), ("Heading 1", 14, 16, 7), ("Heading 2", 11.5, 10, 5)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = FONT, Pt(size), True
        st.font.color.rgb = RGBColor(20, 42, 60)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    # Explicit East Asian font on styles for Word rendering.
    for st in (normal, doc.styles["Title"], doc.styles["Heading 1"], doc.styles["Heading 2"]):
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:eastAsia"), FONT)
    # Word's built-in Title can carry a blue bottom rule through its style.
    title_ppr = doc.styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("大模型产业研发中的智能体应用"), 8, color=(110, 119, 127))
    i = 0
    in_references = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            chunk = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                chunk.append(lines[i])
                i += 1
            add_table(doc, parse_rows(chunk))
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.paragraph_format.first_line_indent = Cm(0)
            ppr = p._p.get_or_add_pPr()
            for border in ppr.findall(qn("w:pBdr")):
                ppr.remove(border)
            add_inline(p, line[2:], 19)
        elif line.startswith("## "):
            in_references = line == "## 参考资料"
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, line[3:], 14)
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, line[4:], 11.5)
        elif line.startswith("!["):
            match = re.match(r"!\[(.+?)\]\((.+?)\)", line)
            if not match:
                raise ValueError(f"Invalid image syntax: {line}")
            caption, name = match.groups()
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(ROOT / name), width=Cm(15.4))
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.first_line_indent = Cm(0)
            cp.paragraph_format.space_after = Pt(9)
            font(cp.add_run(caption), 8.5, color=(86, 99, 111))
        elif line.startswith("- "):
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.45)
            if in_references:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                add_inline(p, "• " + line[2:].replace("*", ""), 8.8)
            else:
                add_inline(p, "• " + line[2:])
        else:
            p = doc.add_paragraph(style="Normal")
            add_inline(p, line.replace("  ", " "))
        i += 1
    doc.core_properties.title = "智能体在大模型产业研发中的应用"
    doc.core_properties.subject = "大模型与智能体工程实践课程报告"
    doc.save(DOCX)


def export_pdf():
    import win32com.client

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = None
    try:
        doc = word.Documents.Open(str(DOCX.resolve()), ReadOnly=True)
        doc.ExportAsFixedFormat(str(PDF.resolve()), 17)
    finally:
        if doc is not None:
            doc.Close(False)
        word.Quit()


def verify_and_render():
    QA.mkdir(parents=True, exist_ok=True)
    pdf = fitz.open(PDF)
    for old in QA.glob("page_*.png"):
        old.unlink()
    findings = []
    for i, page in enumerate(pdf):
        text = page.get_text()
        png = QA / f"page_{i + 1:02d}.png"
        page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3), alpha=False).save(png)
        blocks = page.get_text("blocks")
        outside = [(b[:4], b[4][:45]) for b in blocks if b[0] < 20 or b[2] > page.rect.width - 20 or b[1] < 20 or b[3] > page.rect.height - 20]
        findings.append(f"page {i + 1:02d}: chars={len(text)}, blocks={len(blocks)}, margin_flags={len(outside)}")
    thumbs = []
    for i in range(len(pdf)):
        page_image = Image.open(QA / f"page_{i + 1:02d}.png").convert("RGB")
        page_image.thumbnail((420, 594))
        thumbs.append(page_image)
    sheet = Image.new("RGB", (1260, ((len(thumbs) + 2) // 3) * 594), "#e5e5e5")
    for i, thumb in enumerate(thumbs):
        sheet.paste(thumb, ((i % 3) * 420, (i // 3) * 594))
    sheet.save(QA / "contact_sheet.png")
    all_text = "\n".join(p.get_text() for p in pdf)
    for needle in ("AlphaEvolve", "AgentEvolver", "DGM", "23.3%", "26.4%", "65%", "Runme", "个人分工"):
        if needle not in all_text:
            raise RuntimeError(f"Missing PDF text: {needle}")
    (QA / "render_check.txt").write_text("\n".join(findings) + "\n", encoding="utf-8")
    print(f"DOCX={DOCX}\nPDF={PDF}\nPAGES={len(pdf)}")
    print("\n".join(findings))


if __name__ == "__main__":
    build_docx()
    export_pdf()
    verify_and_render()
