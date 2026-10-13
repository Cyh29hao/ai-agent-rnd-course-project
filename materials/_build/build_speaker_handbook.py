# -*- coding: utf-8 -*-
"""Printable script and cue cards; retain unassigned roles and human pending state."""
from pathlib import Path
import re,json,hashlib,datetime
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import fitz
from PIL import Image
import win32com.client

ROOT=Path(__file__).resolve().parents[1]
SOURCES=['01_四人逐字稿.md','06_四人上场卡.md']
DOCX=ROOT/'12_四人讲稿与上场卡.docx'
PDF=ROOT/'12_四人讲稿与上场卡.pdf'
QA=ROOT/'qa/speaker_handbook'
QA.mkdir(parents=True,exist_ok=True)
doc=Document()
section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=section.bottom_margin=Cm(1.9)
section.left_margin=section.right_margin=Cm(2)
for style in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[style]
    s.font.name='Microsoft YaHei'
    s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.07
doc.styles['Normal'].paragraph_format.space_after=Pt(3)
for name,size in [('Title',23),('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].font.color.rgb=RGBColor.from_string('1A3B4C')
    doc.styles[name].paragraph_format.keep_with_next=True
header=section.header.paragraphs[0]
header.text='备课稿｜A/B/C/D角色待分配；成员实操与排练待完成'
header.runs[0].font.size=Pt(8)
footer=section.footer.paragraphs[0]
footer.alignment=2
footer.add_run('智能体参与大模型产业研发  ·  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(8)

def plain(text):
    text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1（\2）',text)
    return text.replace('**','').replace('`','')

def table(lines):
    rows=[[plain(c.strip()) for c in line.strip().strip('|').split('|')] for line in lines if not re.fullmatch(r'[\s|:\-]+',line)]
    t=doc.add_table(rows=0,cols=max(len(row) for row in rows))
    t.style='Table Grid'
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        prop=cells[0]._tc.getparent().get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
        if i==0:prop.append(OxmlElement('w:tblHeader'))
        for cell,value in zip(cells,row):
            cell.text=value
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                for run in p.runs:run.font.size=Pt(9);run.bold=(i==0)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF0');cell._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph()

for source_index,name in enumerate(SOURCES):
    if source_index:doc.add_page_break()
    lines=(ROOT/name).read_text(encoding='utf-8').splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('|'):
            group=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i].strip());i+=1
            table(group);continue
        if line.startswith('# '):doc.add_paragraph(plain(line[2:]),'Title');continue
        if line.startswith('## '):
            p=doc.add_paragraph(plain(line[3:]),'Heading 1')
            if re.match(r'[ABCD]｜',line[3:]) or line[3:].startswith(('约9分钟','排练后必须补的记录')):p.paragraph_format.page_break_before=True
            continue
        if line.startswith('### '):doc.add_paragraph(plain(line[4:]),'Heading 2');continue
        if line.startswith('> '):
            p=doc.add_paragraph(plain(line[2:]))
            p.paragraph_format.left_indent=Cm(.4)
            p.paragraph_format.space_before=Pt(2)
            for run in p.runs:run.font.size=Pt(11);run.font.color.rgb=RGBColor.from_string('175E50')
            continue
        if line.startswith('- '):doc.add_paragraph(plain(line[2:]),'List Bullet');continue
        doc.add_paragraph(plain(line))
for style in doc.styles:
    ppr=style.element.find(qn('w:pPr'))
    if ppr is not None:
        border=ppr.find(qn('w:pBdr'))
        if border is not None:ppr.remove(border)
for paragraph in doc.paragraphs:
    ppr=paragraph._p.find(qn('w:pPr'))
    if ppr is not None:
        border=ppr.find(qn('w:pBdr'))
        if border is not None:ppr.remove(border)
doc.save(DOCX)
word=win32com.client.DispatchEx('Word.Application')
word.Visible=False;word.DisplayAlerts=0
opened=None
try:
    opened=word.Documents.Open(str(DOCX),ReadOnly=True,AddToRecentFiles=False)
    opened.ExportAsFixedFormat(str(PDF),17)
finally:
    if opened is not None:opened.Close(False)
    if word.Documents.Count==0:word.Quit()

normalize=lambda text:re.sub(r'\s+','',text)
content=json.loads((ROOT/'_build/slide_content.json').read_text(encoding='utf-8'))
stored='\n'.join(p.text for p in Document(DOCX).paragraphs)
with fitz.open(PDF) as pdf:
    extracted='\n'.join(page.get_text() for page in pdf)
    assert all(normalize(row['script']) in normalize(stored) and normalize(row['script']) in normalize(extracted) for row in content)
    assert '\ufffd' not in extracted
    for old in QA.glob('page_*.png'):old.unlink()
    thumbs=[]
    for i,page in enumerate(pdf):
        page.get_pixmap(matrix=fitz.Matrix(1.2,1.2),alpha=False).save(QA/f'page_{i+1:02d}.png')
        im=Image.open(QA/f'page_{i+1:02d}.png').convert('RGB');im.thumbnail((360,510));thumbs.append(im)
    sheet=Image.new('RGB',(1080,510*((len(thumbs)+2)//3)),'#e5e5e5')
    for i,im in enumerate(thumbs):sheet.paste(im,(i%3*360,i//3*510))
    sheet.save(QA/'contact_sheet.png')
    count=len(pdf)
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
manifest={'built_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'sources':{name:sha(ROOT/name) for name in SOURCES},'docx_sha256':sha(DOCX),'pdf_sha256':sha(PDF),'pages':count,'all_14_spoken_passages_stored_and_exported':True,'visual_review':'pending','human_rehearsal':'pending'}
(QA/'build_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'pages':count,'all_14_spoken_passages_stored_and_exported':True,'visual_review':'pending'},ensure_ascii=True))
