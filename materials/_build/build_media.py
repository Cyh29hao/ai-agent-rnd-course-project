"""Build explanatory and log-replay media from checked evidence; no synthetic run records."""
from __future__ import annotations

import json
from pathlib import Path
import time

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
MEDIA = ROOT / "media"
SMOKE = ROOT / "practice" / "runs" / "smoke"
FONT = "Microsoft YaHei"
NAVY = RGBColor(20, 39, 62)
MUTED = RGBColor(79, 102, 123)
BLUE = RGBColor(30, 105, 175)
TEAL = RGBColor(17, 134, 127)
GOLD = RGBColor(191, 119, 31)
PALE = RGBColor(246, 249, 252)


def tx(slide, x, y, w, h, value, size=24, color=NAVY, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    para = frame.paragraphs[0]
    para.text = value
    if align is not None:
        para.alignment = align
    for run in para.runs:
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background() if line is None else None
    return sh


def base(prs, title, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PALE
    rect(slide, 0, 0, 0.14, 7.5, TEAL)
    tx(slide, .55, .42, 12.2, .58, title, 30, NAVY, True)
    tx(slide, .58, 7.02, 12.2, .28, footer, 13, MUTED)
    return slide


def lane(prs, stage):
    slide = base(prs, "一次研究循环：方案经过训练与评分", "据 Anthropic AAR §2–3 重绘｜方法示意，不是公司运行录像")
    names = ["研究代理", "工具与训练", "隔离评测", "共享记录"]
    descriptions = ["提案 mini-paper\n训练代码", "审查代码 → 训练\n产生模型检查点", "运行安全基准\n返回评分与门槛", "保存方法、结果\n供下一轮读取"]
    for i, (name, desc) in enumerate(zip(names, descriptions)):
        yy = 1.55 + i * 1.24
        rect(slide, .58, yy, 2.05, .86, RGBColor(225, 236, 244))
        tx(slide, .78, yy + .22, 1.75, .36, name, 21, NAVY, True)
        color = TEAL if i <= stage else RGBColor(161, 178, 190)
        rect(slide, 3.05, yy, .09, .86, color)
        tx(slide, 3.35, yy + .1, 8.6, .7, desc, 23, color, i == stage)
        if i < 3:
            tx(slide, 11.9, yy + .96, .5, .3, "↓", 24, color)
    tx(slide, .62, 6.56, 11.8, .35,
       "人的目标与实验权限在循环外" if stage == 3 else "这一帧跟踪方案进入训练和评测", 18, GOLD, True)
    return slide


def logs(prs, index, result):
    titles = ["资格探测 ① 读取任务与协议", "资格探测 ② 回查官方来源", "资格探测 ③ 输出并校验 JSON"]
    slide = base(prs, titles[index], "真实公开动作日志重排｜非 UI 录屏｜非按实际运行速度播放")
    steps = [
        [("读取", "题目、来源索引、基线提示"), ("遇到", "先按错误路径找 schema.json"), ("修正", "定位到 practice/schema.json")],
        [("工具", "打开 OpenAI 官方文章"), ("位置", "Running evaluations with Codex"), ("证据", "Codex 先写计划，作者审阅批准后执行")],
        [("产物", f"{result['claim_id']} → {result['label']}"),
         ("定位", f"{result['source_id']}#{result['fragment_id']}"),
         ("校验", "VALID smoke-01 supported T1#T1-run-evals-plan")]
    ][index]
    for i, (kind, body) in enumerate(steps):
        yy = 1.64 + 1.55 * i
        rect(slide, .64, yy, 1.65, .68, [BLUE, GOLD, TEAL][i])
        tx(slide, .78, yy + .13, 1.34, .4, kind, 20, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER)
        tx(slide, 2.62, yy + .09, 9.6, .85, body, 24, NAVY, i == 2)
    tx(slide, .66, 6.42, 12, .42,
       "制作代理资格探测；不是四名学生的已完成实践" if index == 2 else "动作有事后补记；早期步骤的精确时间不可得",
       18, GOLD, True)
    return slide


def inspect_smoke():
    events = [json.loads(line) for line in (SMOKE / "actions.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()]
    result = json.loads((SMOKE / "output.json").read_text(encoding="utf-8"))
    assert result["claim_id"] == "smoke-01" and result["label"] == "supported"
    assert result["source_id"] == "T1" and result["fragment_id"] == "T1-run-evals-plan"
    assert any("schema.json initially requested at root path and missing" in e.get("result", "") for e in events)
    assert any("located and read schema" in e.get("result", "") for e in events)
    assert any(e.get("tool") == "web__run" and "official page" in e.get("result", "") for e in events)
    assert any("wrote JSON decision" in e.get("result", "") for e in events)
    assert any("VALID smoke-01 supported T1#T1-run-evals-plan" in e.get("result", "") for e in events)
    return result


def save_deck(name, seconds, builder):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for i in range(seconds // 5 if name == "anthropic_cycle" else seconds // 6):
        builder(prs, i)
    path = MEDIA / f"{name}_frames.pptx"
    prs.save(path)
    return path


def export_with_powerpoint(paths):
    import win32com.client
    app = win32com.client.DispatchEx("PowerPoint.Application")
    try:
        for name, path, duration in paths:
            pres = app.Presentations.Open(str(path), ReadOnly=True, Untitled=False, WithWindow=False)
            try:
                for i in range(1, pres.Slides.Count + 1):
                    pres.Slides(i).Export(str(MEDIA / f"{name}_{i:02}.png"), "PNG", 1920, 1080)
                target = MEDIA / f"{name}.mp4"
                pres.CreateVideo(str(target), False, duration, 1080, 25, 85)
                deadline = time.monotonic() + 360
                while time.monotonic() < deadline:
                    state = pres.CreateVideoStatus
                    if state == 3 and target.exists() and target.stat().st_size > 10000:
                        break
                    if state == 4:
                        raise RuntimeError(f"PowerPoint video export failed: {name}")
                    time.sleep(2)
                else:
                    raise TimeoutError(f"PowerPoint video export timed out: {name}; status={state}")
                print(f"{target} size={target.stat().st_size} bytes")
            finally:
                pres.Close()
    finally:
        app.Quit()


def main():
    MEDIA.mkdir(exist_ok=True)
    result = inspect_smoke()
    a = save_deck("anthropic_cycle", 10, lambda prs, i: lane(prs, 1 if i == 0 else 3))
    b = save_deck("smoke_log_replay", 18, lambda prs, i: logs(prs, i, result))
    export_with_powerpoint([("anthropic_cycle", a, 5), ("smoke_log_replay", b, 6)])


if __name__ == "__main__":
    main()
