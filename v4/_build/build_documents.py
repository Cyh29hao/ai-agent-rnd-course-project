"""Create report and printable talk using the bundled python-docx library."""
from pathlib import Path
import re,json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
def setup():
    d=Document();sec=d.sections[0]
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        st=d.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    d.styles['Normal'].font.size=Pt(10.5)
    d.styles['Normal'].paragraph_format.line_spacing=1.2
    d.styles['Normal'].paragraph_format.space_after=Pt(5)
    for name,size in [('Title',19),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
        d.styles[name].font.size=Pt(size)
        d.styles[name].paragraph_format.keep_with_next=True
        pr=d.styles[name].element.get_or_add_pPr()
        for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
        none=OxmlElement('w:pBdr')
        for side in ['top','left','bottom','right','between','bar']:
            e=OxmlElement('w:'+side);e.set(qn('w:val'),'nil');none.append(e)
        pr.append(none)
    f=sec.footer.paragraphs[0];f.alignment=2
    r=f.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.append(field)
    return d
def clean(t):return re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1',t.replace('**','').replace('`',''))
def report():
    d=setup();rows=(ROOT/'05_课程报告.md').read_text(encoding='utf-8').splitlines();i=0
    while i<len(rows):
        t=rows[i].strip();i+=1
        if not t:continue
        if t.startswith('|'):
            table_rows=[t]
            while i<len(rows) and rows[i].lstrip().startswith('|'):table_rows.append(rows[i].strip());i+=1
            vals=[[clean(c.strip()) for c in r.strip('|').split('|')] for r in table_rows if not re.match(r'^\|[\s:|\-]+\|$',r)]
            tab=d.add_table(rows=len(vals),cols=len(vals[0]));tab.autofit=False;tab.style='Table Grid'
            for ri,v in enumerate(vals):
                trPr=tab.rows[ri]._tr.get_or_add_trPr()
                trPr.append(OxmlElement('w:cantSplit'))
                if ri==0:trPr.append(OxmlElement('w:tblHeader'))
                for ci,c in enumerate(v):
                    cell=tab.cell(ri,ci);cell.text=c
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4)
                        for run in para.runs:run.font.size=Pt(9);run.bold=ri==0
                    tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                    for side in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    tcPr.append(borders)
            d.add_paragraph();continue
        img=re.match(r'!\[(.*?)\]\((.*?)\)',t)
        if img:
            d.add_picture(str((ROOT/img.group(2)).resolve()),width=Cm(16))
            d.add_paragraph(clean(img.group(1)));continue
        h=re.match(r'^(#{1,3})\s+(.*)',t)
        if h:
            depth=len(h.group(1));d.add_paragraph(clean(h.group(2)),style='Title' if depth==1 else 'Heading '+str(depth-1));continue
        d.add_paragraph(clean(t))
    d.save(ROOT/'outputs/课程报告_V4.docx')
def talk():
    d=setup();d.add_paragraph('四人汇报口播与快速上场卡',style='Title')
    d.add_paragraph('正文按10分钟设计。时间预算尚未真人试讲。A/B/C/D是可交换的讲述角色，尚未对应姓名。')
    data=json.loads((ROOT/'slide_content.json').read_text(encoding='utf-8'))
    for s in data:
        d.add_paragraph(f'{s["page"]:02}  {s["title"]}',style='Heading 1')
        d.add_paragraph(f'建议角色 {s["speaker"]}　设计预算 {s["seconds"]} 秒')
        d.add_paragraph(s['script'])
        p=d.add_paragraph('本页记住：'+s['key_point']);p.paragraph_format.space_after=Pt(14)
    d.add_paragraph('15分钟扩展与答辩准备',style='Heading 1')
    d.add_paragraph('四段扩展正文、12个备课追问和16题详细答辩，分别见01_自然口播稿、02_四人快速备课和03_答辩手册。此打印稿保留主讲正文，避免上场时从资料堆里找要朗读的句子。')
    d.save(ROOT/'outputs/四人口播_V4.docx')
report();talk();print('DOCX_CREATED report,talk')
